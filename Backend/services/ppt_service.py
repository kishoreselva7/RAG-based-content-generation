import os
import re
import json
import logging
from pathlib import Path
from typing import Dict, Any, List
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import config
from services.retrieval_service import retrieve_top_chunks
from services.generation_service import generate_llm_response

logger = logging.getLogger(__name__)

PPT_PROMPT_TEMPLATE = """
Using ONLY the retrieved syllabus context below, generate structured presentation slides content for topic: "{topic}".
Maximum 12 slides.

Return valid JSON with the following schema:
{{
  "presentation_title": "{topic}",
  "subtitle": "Course Presentation",
  "slides": [
    {{
      "title": "Slide Title",
      "bullets": [
        "Bullet point 1",
        "Bullet point 2",
        "Bullet point 3",
        "Bullet point 4"
      ],
      "example": "Real world example or case study",
      "summary": "Slide key summary"
    }}
  ]
}}

Ensure every slide has Title, exactly 4 Bullet Points, Example, and Summary.
Return ONLY raw valid JSON.
"""

def generate_powerpoint(topic: str) -> Dict[str, Any]:
    """Feature 2: PowerPoint Generation (JSON -> PPTX via python-pptx)."""
    chunks, context = retrieve_top_chunks(topic, top_k=config.TOP_K_RETRIEVAL)
    prompt = PPT_PROMPT_TEMPLATE.format(topic=topic)
    
    json_str = generate_llm_response(prompt=prompt, retrieved_context=context, json_mode=True)
    ppt_data = _parse_ppt_json(json_str, topic)

    topic_slug = re.sub(r'[^a-zA-Z0-9_]', '_', topic.lower())[:40]
    filename = f"presentation_{topic_slug}.pptx"
    file_path = config.PPT_DIR / filename

    _create_pptx_file(ppt_data, str(file_path))

    return {
        "topic": topic,
        "ppt_data": ppt_data,
        "file_path": str(file_path),
        "filename": filename,
        "slides_count": len(ppt_data.get("slides", [])),
        "retrieved_chunks": len(chunks)
    }

def _parse_ppt_json(json_str: str, topic: str) -> Dict[str, Any]:
    """Extracts JSON object from string robustly."""
    try:
        match = re.search(r'\{.*\}', json_str, re.DOTALL)
        if match:
            return json.loads(match.group(0))
        return json.loads(json_str)
    except Exception as e:
        logger.warning(f"Failed to parse PPT JSON: {e}. Building fallback slide structure.")
        return {
            "presentation_title": f"Presentation: {topic}",
            "subtitle": "Generated from Syllabus Context",
            "slides": [
                {
                    "title": f"Introduction to {topic}",
                    "bullets": [
                        "Overview of core course objectives.",
                        "Foundational principles derived from syllabus.",
                        "Key methodology and technical framework.",
                        "Academic outcomes and mastery criteria."
                    ],
                    "example": f"Industry case study applying {topic} concepts.",
                    "summary": "Core foundational understanding established."
                },
                {
                    "title": f"Key Syllabus Concepts - {topic}",
                    "bullets": [
                        "Primary theoretical models and formulations.",
                        "Detailed component breakdown.",
                        "Implementation and evaluation steps.",
                        "Critical analysis of syllabus topics."
                    ],
                    "example": "Practical exercise solving domain specific problem.",
                    "summary": "Comprehensive overview of essential topic areas."
                }
            ]
        }

def _create_pptx_file(ppt_data: Dict[str, Any], output_path: str):
    """Creates a styled PPTX file using python-pptx."""
    prs = Presentation()
    prs.slide_width = Inches(13.33)  # 16:9 widescreen format
    prs.slide_height = Inches(7.5)

    # Color Palette: Deep Blue / Purple modern aesthetic
    DARK_BG = RGBColor(15, 23, 42)      # Slate 900
    TEXT_LIGHT = RGBColor(248, 250, 252) # Slate 50
    ACCENT_CYAN = RGBColor(56, 189, 248) # Sky 400
    ACCENT_PURPLE = RGBColor(168, 85, 247) # Purple 500
    TEXT_MUTED = RGBColor(203, 213, 225) # Slate 300
    BOX_BG = RGBColor(30, 41, 59)       # Slate 800

    blank_slide_layout = prs.slide_layouts[6]

    # --- Slide 1: Title Slide ---
    slide1 = prs.slides.add_slide(blank_slide_layout)
    bg1 = slide1.background
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = DARK_BG

    # Title Box
    title_box = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.33), Inches(2.0))
    tf1 = title_box.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = ppt_data.get("presentation_title", "EduGen AI Presentation")
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = ACCENT_CYAN
    p1.alignment = PP_ALIGN.LEFT

    # Subtitle Box
    sub_p = tf1.add_paragraph()
    sub_p.text = ppt_data.get("subtitle", "Generated using RAG & Uploaded Syllabus Context")
    sub_p.font.size = Pt(22)
    sub_p.font.color.rgb = TEXT_MUTED
    sub_p.alignment = PP_ALIGN.LEFT

    # --- Slide Content Loop ---
    slides_list = ppt_data.get("slides", [])
    for idx, slide_info in enumerate(slides_list[:12], 1):
        slide = prs.slides.add_slide(blank_slide_layout)
        bg = slide.background
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG

        # Slide Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.5), Inches(1.0))
        tf = t_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"Slide {idx}: {slide_info.get('title', 'Topic Overview')}"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CYAN

        # Bullets Container
        b_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(7.2), Inches(5.0))
        btf = b_box.text_frame
        btf.word_wrap = True
        
        bullets = slide_info.get("bullets", [])
        for b_idx, bullet in enumerate(bullets[:4]):
            bp = btf.paragraphs[0] if b_idx == 0 else btf.add_paragraph()
            bp.text = f"• {bullet}"
            bp.font.size = Pt(18)
            bp.font.color.rgb = TEXT_LIGHT
            bp.space_after = Pt(14)

        # Right Side Box: Example & Summary
        ex_shape = slide.shapes.add_shape(
            1, Inches(8.3), Inches(1.6), Inches(4.2), Inches(2.5) # MSO_SHAPE.RECTANGLE = 1
        )
        ex_shape.fill.solid()
        ex_shape.fill.fore_color.rgb = BOX_BG
        ex_shape.line.color.rgb = ACCENT_PURPLE
        ex_tf = ex_shape.text_frame
        ex_tf.word_wrap = True
        ex_p1 = ex_tf.paragraphs[0]
        ex_p1.text = "💡 Real-World Example"
        ex_p1.font.bold = True
        ex_p1.font.size = Pt(16)
        ex_p1.font.color.rgb = ACCENT_PURPLE
        ex_p2 = ex_tf.add_paragraph()
        ex_p2.text = slide_info.get("example", "Practical application of concepts.")
        ex_p2.font.size = Pt(14)
        ex_p2.font.color.rgb = TEXT_MUTED

        sum_shape = slide.shapes.add_shape(
            1, Inches(8.3), Inches(4.3), Inches(4.2), Inches(2.3)
        )
        sum_shape.fill.solid()
        sum_shape.fill.fore_color.rgb = BOX_BG
        sum_shape.line.color.rgb = ACCENT_CYAN
        sum_tf = sum_shape.text_frame
        sum_tf.word_wrap = True
        sum_p1 = sum_tf.paragraphs[0]
        sum_p1.text = "📌 Slide Summary"
        sum_p1.font.bold = True
        sum_p1.font.size = Pt(16)
        sum_p1.font.color.rgb = ACCENT_CYAN
        sum_p2 = sum_tf.add_paragraph()
        sum_p2.text = slide_info.get("summary", "Key takeaway point.")
        sum_p2.font.size = Pt(14)
        sum_p2.font.color.rgb = TEXT_MUTED

    prs.save(output_path)
